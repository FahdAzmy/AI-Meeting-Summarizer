"""
src/helpers/excel_generator.py
-------------------------------
Pure utility module for generating Excel (.xlsx) files from Meeting documents.

Uses openpyxl to create styled workbooks in-memory (BytesIO) for streaming
via FastAPI's StreamingResponse.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1C1917", end_color="1C1917", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

COLUMNS = ["Title", "Date", "Duration (min)", "Participants", "Summary", "Decisions", "Action Items", "Follow-up"]


def _get_display_title(meeting: Any) -> str:
    if meeting.title:
        return meeting.title
    platform = meeting.platform or "Meeting"
    date_str = meeting.created_at.strftime("%b %d, %Y")
    return f"{platform} — {date_str}"


def _get_participants(meeting: Any) -> str:
    stats = meeting.speaker_stats
    if not stats:
        return "N/A"
    speakers = stats.get("speakers", []) if isinstance(stats, dict) else []
    names = [s.get("speaker", "") for s in speakers if s.get("speaker")]
    return ", ".join(names) if names else "N/A"


def _format_action_items(items: list[dict]) -> str:
    lines = []
    for a in items:
        assignee = a.get("assignee", "?")
        task = a.get("task", "?")
        deadline = a.get("deadline", "")
        line = f"{assignee}: {task}"
        if deadline:
            line += f" ({deadline})"
        lines.append(line)
    return "\n".join(lines)


def _meeting_to_row(meeting: Any) -> list[Any]:
    return [
        _get_display_title(meeting),
        meeting.created_at.strftime("%Y-%m-%d"),
        meeting.duration_minutes or 0,
        _get_participants(meeting),
        meeting.summary or "",
        "\n".join(meeting.decisions) if meeting.decisions else "",
        _format_action_items(meeting.action_items) if meeting.action_items else "",
        "\n".join(meeting.follow_up) if meeting.follow_up else "",
    ]


def _style_worksheet(ws: Any) -> None:
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    col_widths = [30, 12, 14, 25, 50, 40, 40, 30]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width


def generate_all_meetings_excel(meetings: list[Any]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Meetings"

    _style_worksheet(ws)

    for meeting in meetings:
        ws.append(_meeting_to_row(meeting))

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_single_meeting_excel(meeting: Any) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Meeting"

    _style_worksheet(ws)
    ws.append(_meeting_to_row(meeting))

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
