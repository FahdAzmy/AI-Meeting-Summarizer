"""
tests/unit/test_excel_generator.py
-----------------------------------
TDD test suite for the Excel generation helpers.

Tests are written FIRST (before implementation) following Constitution Principle I.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any
from unittest.mock import MagicMock

import pytest


# ---------------------------------------------------------------------------
# Shared test data
# ---------------------------------------------------------------------------

def _make_mock_meeting(**overrides: Any) -> MagicMock:
    meeting = MagicMock()
    meeting.title = overrides.get("title", "Sprint Standup")
    meeting.platform = overrides.get("platform", "Google Meet")
    meeting.created_at = overrides.get("created_at", MagicMock(strftime=MagicMock(return_value="2026-04-26")))
    meeting.duration_minutes = overrides.get("duration_minutes", 45)
    meeting.summary = overrides.get("summary", "Team discussed sprint goals.")
    meeting.decisions = overrides.get("decisions", ["Approved Q3 plan"])
    meeting.action_items = overrides.get(
        "action_items",
        [{"assignee": "Alice", "task": "Write spec", "deadline": "Friday"}],
    )
    meeting.follow_up = overrides.get("follow_up", ["Schedule retro"])
    meeting.speaker_stats = overrides.get("speaker_stats", {
        "speakers": [
            {"speaker": "Alice", "total_speaking_time_sec": 120},
            {"speaker": "Bob", "total_speaking_time_sec": 80},
        ]
    })
    meeting.transcript = overrides.get("transcript", "Alice: Let's go. Bob: Agreed.")
    return meeting


# ---------------------------------------------------------------------------
# T010 — Excel generator: generate_all_meetings_excel
# ---------------------------------------------------------------------------

class TestGenerateAllMeetingsExcel:
    """T010: generate_all_meetings_excel must return a valid xlsx BytesIO buffer."""

    def test_returns_bytes_io(self) -> None:
        from src.helpers.excel_generator import generate_all_meetings_excel
        meetings = [_make_mock_meeting()]
        result = generate_all_meetings_excel(meetings)
        assert isinstance(result, BytesIO)

    def test_xlsx_is_valid_workbook(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_all_meetings_excel
        meetings = [_make_mock_meeting(), _make_mock_meeting(title="Retro")]
        buf = generate_all_meetings_excel(meetings)
        wb = load_workbook(buf)
        ws = wb.active
        assert ws is not None

    def test_header_row_contains_expected_columns(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_all_meetings_excel
        meetings = [_make_mock_meeting()]
        buf = generate_all_meetings_excel(meetings)
        wb = load_workbook(buf)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        for col in ["Title", "Date", "Duration (min)", "Participants", "Summary", "Decisions", "Action Items", "Follow-up"]:
            assert col in headers

    def test_data_row_count_matches_meetings(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_all_meetings_excel
        meetings = [_make_mock_meeting(), _make_mock_meeting(title="Retro")]
        buf = generate_all_meetings_excel(meetings)
        wb = load_workbook(buf)
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 2

    def test_meeting_title_in_data_row(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_all_meetings_excel
        meetings = [_make_mock_meeting(title="My Meeting")]
        buf = generate_all_meetings_excel(meetings)
        wb = load_workbook(buf)
        ws = wb.active
        first_row = [cell.value for cell in ws[2]]
        assert "My Meeting" in first_row

    def test_empty_meetings_returns_header_only(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_all_meetings_excel
        buf = generate_all_meetings_excel([])
        wb = load_workbook(buf)
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 0

    def test_title_fallback_when_none(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_all_meetings_excel
        meeting = _make_mock_meeting(title=None, platform="Zoom")
        meeting.created_at.strftime = MagicMock(return_value="Apr 26, 2026")
        buf = generate_all_meetings_excel([meeting])
        wb = load_workbook(buf)
        ws = wb.active
        first_row = [cell.value for cell in ws[2]]
        assert any("Zoom" in str(v) for v in first_row if v)

    def test_participants_extracted_from_speaker_stats(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_all_meetings_excel
        meetings = [_make_mock_meeting()]
        buf = generate_all_meetings_excel(meetings)
        wb = load_workbook(buf)
        ws = wb.active
        row_values = [str(cell.value or "") for cell in ws[2]]
        combined = " ".join(row_values)
        assert "Alice" in combined
        assert "Bob" in combined


# ---------------------------------------------------------------------------
# T022 — Excel generator: generate_single_meeting_excel
# ---------------------------------------------------------------------------

class TestGenerateSingleMeetingExcel:
    """T022: generate_single_meeting_excel must return a valid xlsx BytesIO with one data row."""

    def test_returns_bytes_io(self) -> None:
        from src.helpers.excel_generator import generate_single_meeting_excel
        meeting = _make_mock_meeting()
        result = generate_single_meeting_excel(meeting)
        assert isinstance(result, BytesIO)

    def test_xlsx_has_one_data_row(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_single_meeting_excel
        meeting = _make_mock_meeting(title="Single Export")
        buf = generate_single_meeting_excel(meeting)
        wb = load_workbook(buf)
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 1

    def test_header_matches_all_meetings_format(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_single_meeting_excel
        meeting = _make_mock_meeting()
        buf = generate_single_meeting_excel(meeting)
        wb = load_workbook(buf)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        for col in ["Title", "Date", "Duration (min)", "Participants", "Summary", "Decisions", "Action Items", "Follow-up"]:
            assert col in headers

    def test_meeting_title_in_row(self) -> None:
        from openpyxl import load_workbook
        from src.helpers.excel_generator import generate_single_meeting_excel
        meeting = _make_mock_meeting(title="Budget Review")
        buf = generate_single_meeting_excel(meeting)
        wb = load_workbook(buf)
        ws = wb.active
        first_row = [cell.value for cell in ws[2]]
        assert "Budget Review" in first_row
